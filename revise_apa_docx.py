from pathlib import Path

from docx import Document
from docx.shared import Inches


DOCX = Path(r"C:\Users\TomTom Catral\Documents\blgfdts\MSIT-FINAL-FINAL-APA-REVISED.docx")


def replace_with_citation(paragraph, citation: str) -> None:
    text = paragraph.text.rstrip()
    if citation in text:
        return
    if text.endswith("."):
        text = text[:-1]
    paragraph.clear()
    paragraph.add_run(f"{text} {citation}.")


def write_reference(paragraph, segments) -> None:
    paragraph.clear()
    for text, italic in segments:
        run = paragraph.add_run(text)
        run.italic = italic
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.5)
    fmt.first_line_indent = Inches(-0.5)
    fmt.line_spacing = 2
    fmt.space_before = 0
    fmt.space_after = 0


doc = Document(DOCX)

# APA 7 author-date citations placed where the corresponding borrowed ideas occur.
citations = {
    89: "(International Organization for Standardization [ISO], 2011)",
    96: "(International Organization for Standardization [ISO], 2016)",
    138: "(ISO, 2016)",
    142: "(ISO, 2011)",
    144: "(Republic Act No. 8792, 2000; Republic Act No. 10173, 2012)",
    145: "(Republic Act No. 11032, 2018)",
    146: "(National Archives of the Philippines, n.d.)",
    153: "(Pressman & Maxim, 2020; Sommerville, 2016)",
    155: "(Express.js, n.d.; OpenJS Foundation, n.d.; Oracle Corporation, n.d.; React Team, n.d.)",
    171: "(ISO, 2011)",
}
for index, citation in citations.items():
    replace_with_citation(doc.paragraphs[index], citation)

# APA calls the source list “References.” Keep the existing heading style and position.
doc.paragraphs[558].text = "REFERENCES"

references = [
    [("Express.js. (n.d.). ", False), ("Express: Fast, unopinionated, minimalist web framework for Node.js", True), (". https://expressjs.com/", False)],
    [("International Organization for Standardization. (2011). ", False), ("Systems and software engineering—Systems and software Quality Requirements and Evaluation (SQuaRE)—System and software quality models", True), (" (ISO/IEC Standard No. 25010:2011). https://www.iso.org/standard/35733.html", False)],
    [("International Organization for Standardization. (2016). ", False), ("Information and documentation—Records management—Part 1: Concepts and principles", True), (" (ISO Standard No. 15489-1:2016). https://www.iso.org/standard/62542.html", False)],
    [("National Archives of the Philippines. (n.d.). ", False), ("Frequently asked questions", True), (". https://nationalarchives.gov.ph/frequently-asked-question", False)],
    [("OpenJS Foundation. (n.d.). ", False), ("Node.js documentation", True), (". https://nodejs.org/docs/latest/api/", False)],
    [("Oracle Corporation. (n.d.). ", False), ("MySQL reference manual", True), (". https://dev.mysql.com/doc/", False)],
    [("Pressman, R. S., & Maxim, B. R. (2020). ", False), ("Software engineering: A practitioner’s approach", True), (" (9th ed.). McGraw Hill.", False)],
    [("React Team. (n.d.). ", False), ("React documentation", True), (". https://react.dev/", False)],
    [("Republic Act No. 8792. (2000). ", False), ("Electronic Commerce Act of 2000", True), (". Official Gazette of the Republic of the Philippines. https://www.officialgazette.gov.ph/2000/06/14/republic-act-no-8792-s-2000/", False)],
    [("Republic Act No. 10173. (2012). ", False), ("Data Privacy Act of 2012", True), (". Official Gazette of the Republic of the Philippines. https://www.officialgazette.gov.ph/2012/08/15/republic-act-no-10173/", False)],
    [("Republic Act No. 11032. (2018). ", False), ("Ease of Doing Business and Efficient Government Service Delivery Act of 2018", True), (". Official Gazette of the Republic of the Philippines. https://www.officialgazette.gov.ph/2018/05/28/republic-act-no-11032/", False)],
    [("Sommerville, I. (2016). ", False), ("Software engineering", True), (" (10th ed.). Pearson.", False)],
]

for paragraph, reference in zip(doc.paragraphs[560:572], references):
    write_reference(paragraph, reference)

doc.save(DOCX)
print(DOCX)
